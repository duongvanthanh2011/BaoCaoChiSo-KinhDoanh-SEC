"""
report_components.py - Component UI dung chung cho cac bao cao.
"""

import hashlib
import io
from contextlib import contextmanager

import pandas as pd
import streamlit as st
from st_aggrid import AgGrid
from st_aggrid.grid_options_builder import GridOptionsBuilder

from manual_input_schema import (
    REPORT_1_CODE,
    REPORT_2_CODE,
    get_widget_key,
    get_r2_duplicate_code,
)
from manual_input_state import (
    on_widget_change,
    get_draft_value,
    set_draft_value,
    get_changed_values,
    has_any_draft_changes,
    get_session_entry,
    sync_manual_widgets,
    apply_saved_snapshot,

    set_loaded_baseline,
)
from time_utils import get_vn_now


DOT_COLUMN = 'ĐỢT HỌC THỬ'


@contextmanager
def manual_input_expander(title, column_spec=None, expanded=True):
    """Tạo khối nhập liệu có thể đóng/mở và tùy chọn bố cục cột."""
    with st.expander(title, expanded=expanded):
        columns = st.columns(column_spec) if column_spec is not None else None
        yield columns


def normalize_dot_manual_df(dot_manual_df, unique_dots, number_columns, float_columns=None):
    """Giữ đúng danh sách đợt và ép các cột nhập tay về số không âm (hỗ trợ cả float và int)."""
    base_df = pd.DataFrame({DOT_COLUMN: list(unique_dots)})
    columns = [DOT_COLUMN] + list(number_columns)
    float_cols_set = set(float_columns or [])

    if dot_manual_df is None or len(dot_manual_df) == 0:
        for col in number_columns:
            base_df[col] = 0.0 if col in float_cols_set else 0
        return base_df[columns]

    manual_df = pd.DataFrame(dot_manual_df).copy()
    for col in columns:
        if col not in manual_df.columns:
            manual_df[col] = 0.0 if col in float_cols_set else (0 if col in number_columns else "")

    manual_df = manual_df[columns].drop_duplicates(DOT_COLUMN, keep='last')
    normalized_df = base_df.merge(manual_df, on=DOT_COLUMN, how='left')

    for col in number_columns:
        if col in float_cols_set:
            normalized_df[col] = (
                pd.to_numeric(normalized_df[col], errors='coerce')
                .fillna(0.0)
                .clip(lower=0.0)
                .round(2)
            )
        else:
            normalized_df[col] = (
                pd.to_numeric(normalized_df[col], errors='coerce')
                .fillna(0)
                .clip(lower=0)
                .astype(int)
            )

    return normalized_df[columns]


def hash_dot_manual_df(dot_manual_df):
    """Tao hash on dinh de ep component bang bao cao reload khi so nhap tay doi."""
    payload = dot_manual_df.to_json(orient='records', force_ascii=False)
    return hashlib.md5(payload.encode('utf-8')).hexdigest()[:12]


def render_dot_manual_inputs(
    title,
    state_key,
    unique_dots,
    number_columns,
    input_key_prefixes,
    float_columns=None,
    display_labels=None,
    report_code=None,
):
    """
    Render bảng nhập số theo từng đợt học thử với định dạng và chiều cao đồng bộ chuẩn xác.
    Được tích hợp với bản nháp (draft) trong phiên, tự động đồng bộ khi người dùng gõ số.

    Trả về:
    - manual_df: DataFrame gồm DOT_COLUMN và các number_columns.
    - manual_hash: hash của manual_df để dùng làm một phần key cho AgGrid.
    """
    unique_dots = list(unique_dots)
    number_columns = list(number_columns)
    float_cols_set = set(float_columns or [])
    labels_map = display_labels or {}
    
    st.session_state[state_key] = normalize_dot_manual_df(
        st.session_state.get(state_key),
        unique_dots,
        number_columns,
        float_columns=float_columns
    )

    # Tiêu đề với chiều cao và font size đồng nhất
    st.markdown(f"<div style='font-size: 15px; font-weight: 600; color: #1E293B; margin-bottom: 8px; height: 24px; line-height: 24px;'>{title}</div>", unsafe_allow_html=True)

    manual_by_dot = st.session_state[state_key].set_index(DOT_COLUMN)
    col_widths = [1.3] + [1.0] * len(number_columns)
    header_cols = st.columns(col_widths)
    header_cols[0].markdown(f"<div style='height: 38px; display: flex; align-items: flex-end; font-weight: 700; font-size: 13px;'>{DOT_COLUMN}</div>", unsafe_allow_html=True)
    for idx, col in enumerate(number_columns, start=1):
        col_label = labels_map.get(col, col)
        header_cols[idx].markdown(f"<div style='height: 38px; display: flex; align-items: flex-end; font-weight: 700; font-size: 13px; line-height: 1.2;'>{col_label}</div>", unsafe_allow_html=True)

    manual_rows = []
    for dot_name in unique_dots:
        dot_hash = hashlib.md5(str(dot_name).encode("utf-8")).hexdigest()[:12]
        row_cols = st.columns(col_widths)
        row_cols[0].markdown(f"<div style='height: 38px; display: flex; align-items: center; font-size: 13px; font-weight: 500;'>{dot_name}</div>", unsafe_allow_html=True)

        row_data = {DOT_COLUMN: dot_name}
        for idx, col in enumerate(number_columns, start=1):
            key_prefix = input_key_prefixes.get(col, f"{state_key}_{idx}")
            input_code = key_prefix
            is_float = col in float_cols_set

            if report_code:
                input_key = get_widget_key(report_code, input_code, str(dot_name))
                draft_val = get_draft_value(report_code, str(dot_name), input_code, 0)
                if input_key not in st.session_state:
                    st.session_state[input_key] = float(draft_val) if is_float else int(draft_val)
            else:
                input_key = f"{key_prefix}_{dot_hash}"
                if is_float:
                    current_value = float(manual_by_dot.at[dot_name, col]) if dot_name in manual_by_dot.index else 0.0
                    if input_key not in st.session_state:
                        st.session_state[input_key] = current_value
                else:
                    current_value = int(manual_by_dot.at[dot_name, col]) if dot_name in manual_by_dot.index else 0
                    if input_key not in st.session_state:
                        st.session_state[input_key] = current_value

            on_change_kw = {}
            if report_code:
                on_change_kw = {
                    "on_change": on_widget_change,
                    "args": (report_code, str(dot_name), input_code, input_key),
                }

            if is_float:
                val = row_cols[idx].number_input(
                    col,
                    min_value=0.0,
                    step=0.1,
                    format="%.2f",
                    key=input_key,
                    label_visibility="collapsed",
                    **on_change_kw,
                )
                row_data[col] = val
                if report_code:
                    set_draft_value(report_code, str(dot_name), input_code, int(val))
            else:
                val = row_cols[idx].number_input(
                    col,
                    min_value=0,
                    step=1,
                    format="%d",
                    key=input_key,
                    label_visibility="collapsed",
                    **on_change_kw,
                )
                row_data[col] = val
                if report_code:
                    set_draft_value(report_code, str(dot_name), input_code, int(val))

        manual_rows.append(row_data)

    manual_df = normalize_dot_manual_df(pd.DataFrame(manual_rows), unique_dots, number_columns, float_columns=float_columns)
    st.session_state[state_key] = manual_df
    return manual_df, hash_dot_manual_df(manual_df)


def normalize_dot_nguon_manual_df(dot_manual_df, unique_pairs, number_columns):
    """Giữ đúng danh sách (đợt, nguồn) và ép các cột nhập tay về số nguyên không âm."""
    base_df = pd.DataFrame(unique_pairs, columns=[DOT_COLUMN, 'Nguồn'])
    columns = [DOT_COLUMN, 'Nguồn'] + list(number_columns)

    if dot_manual_df is None or len(dot_manual_df) == 0:
        for col in number_columns:
            base_df[col] = 0
        return base_df[columns]

    manual_df = pd.DataFrame(dot_manual_df).copy()
    for col in columns:
        if col not in manual_df.columns:
            manual_df[col] = 0 if col in number_columns else ""

    manual_df = manual_df[columns].drop_duplicates([DOT_COLUMN, 'Nguồn'], keep='last')
    normalized_df = base_df.merge(manual_df, on=[DOT_COLUMN, 'Nguồn'], how='left')

    for col in number_columns:
        normalized_df[col] = (
            pd.to_numeric(normalized_df[col], errors='coerce')
            .fillna(0)
            .clip(lower=0)
            .astype(int)
        )

    return normalized_df[columns]


def render_dot_nguon_manual_inputs(title, state_key, unique_pairs, number_columns, input_key_prefixes):
    """
    Render bảng nhập số theo từng đợt học thử VÀ nguồn.
    unique_pairs là danh sách các tuple (ĐỢT HỌC THỬ, Nguồn).
    """
    number_columns = list(number_columns)
    st.session_state[state_key] = normalize_dot_nguon_manual_df(
        st.session_state.get(state_key),
        unique_pairs,
        number_columns,
    )

    st.markdown(f"#### {title}")

    # Set index to both DOT_COLUMN and 'Nguồn'
    manual_by_pair = st.session_state[state_key].set_index([DOT_COLUMN, 'Nguồn'])
    
    # 2 columns for headers, plus number_columns
    header_cols = st.columns([1.5, 1.5] + [1] * len(number_columns))
    header_cols[0].markdown(f"**{DOT_COLUMN}**")
    header_cols[1].markdown("**Nguồn**")
    for idx, col in enumerate(number_columns, start=2):
        header_cols[idx].markdown(f"**{col}**")

    manual_rows = []
    # Khai báo biến lưu đợt học thử hiện tại để hiển thị đẹp hơn
    current_dot = None

    for dot_name, nguon_name in unique_pairs:
        pair_hash = hashlib.md5(f"{dot_name}_{nguon_name}".encode("utf-8")).hexdigest()[:12]
        row_cols = st.columns([1.5, 1.5] + [1] * len(number_columns))
        
        # Chỉ in tên đợt ở dòng đầu tiên của đợt đó
        if dot_name != current_dot:
            row_cols[0].write(f"**{dot_name}**")
            current_dot = dot_name
        else:
            row_cols[0].write("")

        row_cols[1].write(str(nguon_name))

        row_data = {DOT_COLUMN: dot_name, 'Nguồn': nguon_name}
        for idx, col in enumerate(number_columns, start=2):
            key_prefix = input_key_prefixes.get(col, f"{state_key}_{idx}")
            input_key = f"{key_prefix}_{pair_hash}"
            
            current_value = 0
            if (dot_name, nguon_name) in manual_by_pair.index:
                current_value = int(manual_by_pair.at[(dot_name, nguon_name), col])

            if input_key not in st.session_state:
                st.session_state[input_key] = current_value

            row_data[col] = row_cols[idx].number_input(
                col,
                min_value=0,
                step=1,
                format="%d",
                key=input_key,
                label_visibility="collapsed",
            )

        manual_rows.append(row_data)

    manual_df = normalize_dot_nguon_manual_df(pd.DataFrame(manual_rows), unique_pairs, number_columns)
    st.session_state[state_key] = manual_df
    return manual_df, hash_dot_manual_df(manual_df)


def render_dot_metric_matrix_inputs(
    title,
    state_key,
    unique_dots,
    column_labels,
    report_code,
    input_code_factory,
    value_column,
    key_prefix="metric_matrix",
    step=1,
    min_value=0,
    number_format="%d",
):
    """
    Render bảng nhập metric dạng ma trận tổng quát (mỗi dòng là 1 Đợt, các cột là column_labels).
    Tích hợp với bản nháp (draft) trong phiên, tự động đồng bộ khi gõ số.
    
    Trả về:
    - manual_df: DataFrame gồm [DOT_COLUMN, 'Nguồn', value_column]
    - manual_hash: hash của manual_df
    """
    unique_dots = list(unique_dots)
    column_labels = list(column_labels)
    
    if state_key not in st.session_state:
        st.session_state[state_key] = {}
        
    st.markdown(
        f"<div style='font-size: 15px; font-weight: 600; color: #1E293B; margin-bottom: 8px; height: 24px; line-height: 24px;'>{title}</div>",
        unsafe_allow_html=True,
    )
    
    col_widths = [1.3] + [1.0] * len(column_labels)
    header_cols = st.columns(col_widths)
    header_cols[0].markdown(
        f"<div style='height: 38px; display: flex; align-items: flex-end; font-weight: 700; font-size: 13px;'>{DOT_COLUMN}</div>",
        unsafe_allow_html=True,
    )
    for idx, col_lbl in enumerate(column_labels, start=1):
        header_cols[idx].markdown(
            f"<div style='height: 38px; display: flex; align-items: flex-end; font-weight: 700; font-size: 13px; line-height: 1.2;'>{col_lbl}</div>",
            unsafe_allow_html=True,
        )
        
    manual_rows = []
    for dot_name in unique_dots:
        dot_hash = hashlib.md5(str(dot_name).encode("utf-8")).hexdigest()[:8]
        row_cols = st.columns(col_widths)
        row_cols[0].markdown(
            f"<div style='height: 38px; display: flex; align-items: center; font-size: 13px; font-weight: 500;'>{dot_name}</div>",
            unsafe_allow_html=True,
        )
        
        for idx, col_lbl in enumerate(column_labels, start=1):
            if report_code and input_code_factory:
                input_code = input_code_factory(col_lbl)
                input_key = get_widget_key(report_code, input_code, str(dot_name))
                draft_val = get_draft_value(report_code, str(dot_name), input_code, 0)
                if input_key not in st.session_state:
                    st.session_state[input_key] = int(draft_val)
                on_change_kw = {
                    "on_change": on_widget_change,
                    "args": (report_code, str(dot_name), input_code, input_key),
                }
            else:
                col_hash = hashlib.md5(str(col_lbl).encode("utf-8")).hexdigest()[:6]
                input_key = f"{key_prefix}_{dot_hash}_{col_hash}"
                current_val = st.session_state.get(state_key, {}).get((str(dot_name), str(col_lbl)), 0)
                if input_key not in st.session_state:
                    st.session_state[input_key] = int(current_val)
                on_change_kw = {}
                
            val = row_cols[idx].number_input(
                f"{dot_name}_{col_lbl}_{key_prefix}",
                min_value=min_value,
                step=step,
                format=number_format,
                key=input_key,
                label_visibility="collapsed",
                **on_change_kw,
            )
            
            if report_code and input_code_factory:
                set_draft_value(report_code, str(dot_name), input_code, int(val))

            st.session_state[state_key][(str(dot_name), str(col_lbl))] = val
            manual_rows.append({
                DOT_COLUMN: dot_name,
                'Nguồn': col_lbl,
                value_column: val,
            })
            
    manual_df = pd.DataFrame(manual_rows)
    return manual_df, hash_dot_manual_df(manual_df)


def render_dot_nguon_matrix_inputs(title, state_key, unique_dots, unique_nguons, key_prefix="r2_trung", report_code=REPORT_2_CODE):
    """
    Render bảng nhập Data trùng dạng ma trận (mỗi dòng là 1 Đợt, các cột là các Nguồn).
    Tích hợp với bản nháp (draft) trong phiên, tự động đồng bộ khi gõ số.
    Wrapper tương thích cho Báo cáo 2.
    """
    return render_dot_metric_matrix_inputs(
        title=title,
        state_key=state_key,
        unique_dots=unique_dots,
        column_labels=unique_nguons,
        report_code=report_code,
        input_code_factory=get_r2_duplicate_code,
        value_column="Data trùng",
        key_prefix=key_prefix,
    )



# ==========================================
# CÁC COMPONENT HÀNH ĐỘNG: LƯU VÀ TẢI LẠI
# ==========================================

def reload_manual_inputs(report_code, dots, displayed_codes_by_dot, repository):
    """Fetch hết trước khi thay trạng thái; gọi trước các input của báo cáo."""
    latest = repository.load_latest_batch(report_code, dots)
    for dot in dots:
        entry = get_session_entry(report_code, dot)
        # Ghi nhớ key cũ để reset cả key không còn trong server.
        codes = set(displayed_codes_by_dot.get(dot, [])) | set(entry["draft"]) | set(entry["baseline"])
        rec = latest.get(dot, {})
        set_loaded_baseline(report_code, dot, rec.get("version_no", 0),
                            rec.get("metric_value", {}), discard_edits=True)
        sync_manual_widgets(report_code, dot, codes)


def render_report_actions_bar(report_code, unique_dots, displayed_codes_by_dot, repository=None):
    """Nút lưu và tải lại đặt trước input để đồng bộ widget an toàn."""
    if not unique_dots:
        return
    dots = list(dict.fromkeys(str(d) for d in unique_dots))
    scope = tuple(sorted(dots))
    confirm_key = f"confirm_reload_{report_code}"
    notice_key = f"manual_notice_{report_code}"
    for level, message in st.session_state.pop(notice_key, []):
        getattr(st, level)(message)
    # Không xác nhận nhầm phạm vi nếu người dùng đổi bộ lọc giữa chừng.
    if st.session_state.get(confirm_key) != scope:
        st.session_state.pop(confirm_key, None)

    configured = bool(repository and repository.is_configured())
    loaded = all(get_session_entry(report_code, d)["load_status"] == "loaded" for d in dots)
    dirty = has_any_draft_changes(report_code, dots)
    col_save, col_reload, col_status = st.columns([2.0, 2.5, 5.5])
    with col_save:
        save_btn = st.button("💾 Lưu dữ liệu", key=f"btn_save_{report_code}", type="primary",
                             use_container_width=True, disabled=not configured or not loaded,
                             help="Lưu các ô đã sửa đang hiển thị; giữ số chưa lưu của nguồn đang ẩn.")
    with col_reload:
        reload_btn = st.button("🔄 Tải lại dữ liệu đã lưu", key=f"btn_reload_{report_code}",
                               use_container_width=True, disabled=not configured)
    with col_status:
        if not configured:
            st.caption("Chưa kết nối lưu trữ; số nhập chỉ ở trong phiên.")
        elif not loaded:
            st.caption("Chưa tải đủ dữ liệu đã lưu. Bấm Tải lại trước khi lưu.")
        elif dirty:
            st.caption("● Có thay đổi chưa lưu, có thể gồm nguồn đang ẩn.")
        else:
            st.caption("✓ Không có thay đổi chưa lưu trong các đợt đang chọn.")

    do_reload = False
    if reload_btn:
        if dirty:
            st.session_state[confirm_key] = scope
        else:
            do_reload = True
    if st.session_state.get(confirm_key) == scope:
        st.warning("Tải lại sẽ thay thế số chưa lưu trong các đợt đang chọn, kể cả nguồn đang ẩn.")
        yes, no = st.columns(2)
        with yes:
            if st.button("✅ Xác nhận tải lại", key=f"btn_confirm_ok_{report_code}"):
                do_reload = True
        with no:
            if st.button("❌ Hủy", key=f"btn_confirm_cancel_{report_code}"):
                st.session_state.pop(confirm_key, None)
                st.rerun()
    if do_reload:
        try:
            reload_manual_inputs(report_code, dots, displayed_codes_by_dot, repository)
        except Exception:
            # Không thay baseline, draft hoặc widget khi request thất bại.
            st.error("Không tải lại được dữ liệu. Số đang nhập vẫn được giữ; hãy kiểm tra kết nối/migration và thử lại.")
        else:
            st.session_state.pop(confirm_key, None)
            st.session_state[notice_key] = [("success", "Đã tải lại dữ liệu mới nhất.")]
            st.rerun()

    if save_btn and configured and loaded:
        saved, unchanged, failed = [], [], []
        for dot in dots:
            codes = displayed_codes_by_dot.get(dot, [])
            changed = get_changed_values(report_code, dot, codes)
            if not changed:
                unchanged.append(dot)
                continue
            try:
                res = repository.save_changes(report_code, dot, changed)
                # Cả saved=False cũng trả baseline mới nhất từ máy chủ.
                apply_saved_snapshot(report_code, dot, res["version_no"], res["metric_value"], changed)
                sync_manual_widgets(report_code, dot, codes)
                (saved if res["saved"] else unchanged).append(dot)
            except Exception:
                failed.append(dot)
        notices = []
        if saved or unchanged:
            notices.append(("success", f"Đã lưu {len(saved)} đợt; {len(unchanged)} đợt không thay đổi. "
                            f"({get_vn_now().strftime('%H:%M:%S')})"))
        if failed:
            notices.append(("error", "Chưa xác nhận được kết quả lưu cho: " + ", ".join(failed)
                            + ". Bản nháp được giữ; kiểm tra kết nối/migration và đối chiếu dữ liệu trước khi lưu lại."))
        st.session_state[notice_key] = notices
        st.session_state.pop(confirm_key, None)
        st.rerun()


# ==========================================
# CÁC COMPONENT DÙNG CHUNG CHO RENDER BÁO CÁO
# ==========================================

def render_aggrid_report(df, gb, pinned_row=None, grid_key="grid_report", fit_columns=True):
    """
    Render bảng AgGrid chuẩn hóa với các tùy chọn grid options chung cho mọi báo cáo.
    Loại bỏ trùng lặp boilerplate giữa render_report_1, _2, _3.

    Args:
        df: DataFrame hiển thị.
        gb: GridOptionsBuilder đã cấu hình sẵn cột.
        pinned_row: dict hoặc None — dòng tổng cố định ở đầu bảng.
        grid_key: string — key duy nhất cho AgGrid component.
        fit_columns: boolean — tự co giãn cột cho vừa khung nhìn (mặc định True).
    """
    grid_options = gb.build()
    grid_options["groupIncludeFooter"] = True
    grid_options["groupIncludeTotalFooter"] = True
    grid_options["groupDefaultExpanded"] = -1
    grid_options["suppressAggFuncInHeader"] = True

    if pinned_row is not None:
        grid_options["pinnedTopRowData"] = [pinned_row]

    AgGrid(
        df,
        gridOptions=grid_options,
        enable_enterprise_modules=True,
        allow_unsafe_jscode=True,
        fit_columns_on_grid_load=fit_columns,
        height=550,
        server_sync_strategy="server_wins",
        key=grid_key
    )


def get_excel_bytes(df_excel, sheet_name, file_name):
    """Giữ một file gần nhất mỗi báo cáo trong phiên, không chia sẻ dữ liệu giữa người dùng."""
    cache = st.session_state.setdefault('_excel_download_cache', {})
    key = (sheet_name, file_name)
    cached = cache.get(key)
    if cached is not None and cached['frame'].equals(df_excel):
        return cached['data']
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        df_excel.to_excel(writer, sheet_name=sheet_name, index=False)
        for row in writer.sheets[sheet_name].iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = '0.00'
    data = buffer.getvalue()
    cache[key] = {'frame': df_excel.copy(), 'data': data}
    return data


def render_excel_download(df_excel, sheet_name, file_name, button_label):
    """
    Render nút download Excel chuẩn hóa cho mọi báo cáo.
    Loại bỏ trùng lặp boilerplate giữa render_report_1, _2, _3.

    Args:
        df_excel: DataFrame đã chuẩn bị cho xuất Excel.
        sheet_name: Tên sheet trong file Excel.
        file_name: Tên file Excel khi download.
        button_label: Nhãn nút download.
    """
    st.download_button(
        label=button_label,
        data=get_excel_bytes(df_excel, sheet_name, file_name),
        file_name=file_name,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


def build_report_3_excel_bytes(df_excel, file_name):
    """Tạo Excel BC3 có header hai tầng giống bố cục nhóm cột trên giao diện."""
    cache = st.session_state.setdefault('_excel_download_cache', {})
    cached = cache.get(file_name)
    if cached is not None and cached.get('frame') is not None and cached['frame'].equals(df_excel):
        return cached['data']

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from report_calculations import get_report_3_column_specs

    header_fill = PatternFill(start_color='E8F0FE', end_color='E8F0FE', fill_type='solid')
    group_fill = PatternFill(start_color='D2E3FC', end_color='D2E3FC', fill_type='solid')
    bold_font = Font(name='Calibri', size=11, bold=True)
    normal_font = Font(name='Calibri', size=11)
    border = Border(*[Side(style='thin', color='D3D3D3')] * 4)
    specs = get_report_3_column_specs()
    flat_cols = [
        (child_name, field, fmt_type, width)
        for _, children in specs for child_name, field, fmt_type, width in children
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = 'BC_Nguon_Tuoi'
    col_idx = 1
    for group_name, children in specs:
        start_col = col_idx
        end_col = start_col + len(children) - 1
        if group_name is None:
            for child_name, _, _, _ in children:
                cell = ws.cell(1, col_idx, child_name)
                cell.font = bold_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                cell.border = border
                ws.cell(2, col_idx).border = border
                ws.merge_cells(start_row=1, start_column=col_idx, end_row=2, end_column=col_idx)
                col_idx += 1
        else:
            for current_col in range(start_col, end_col + 1):
                ws.cell(1, current_col).fill = group_fill
                ws.cell(1, current_col).border = border
            cell = ws.cell(1, start_col, group_name)
            cell.font = bold_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
            for child_name, _, _, _ in children:
                cell = ws.cell(2, col_idx, child_name)
                cell.font = bold_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                cell.border = border
                col_idx += 1

    for row_idx, (_, row_data) in enumerate(df_excel.iterrows(), start=3):
        is_total = str(row_data.get('ĐỢT HỌC THỬ', '')).startswith('TỔNG')
        for col_idx, (_, field, fmt_type, _) in enumerate(flat_cols, start=1):
            cell = ws.cell(row_idx, col_idx)
            _format_excel_cell(cell, row_data.get(field), fmt_type, is_total, normal_font, bold_font, border)

    for col_idx, (_, _, _, width) in enumerate(flat_cols, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(12, int(width / 7.5))
    ws.freeze_panes = 'D3'
    ws.auto_filter.ref = f"A2:{get_column_letter(len(flat_cols))}{max(2, len(df_excel) + 2)}"

    buffer = io.BytesIO()
    wb.save(buffer)
    data = buffer.getvalue()
    cache[file_name] = {'frame': df_excel.copy(), 'data': data}
    return data


def render_report_3_excel_download(df_excel, file_name, button_label):
    """Nút tải Excel BC3 với header nhóm tuổi hai tầng."""
    st.download_button(
        label=button_label,
        data=build_report_3_excel_bytes(df_excel, file_name),
        file_name=file_name,
        mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


def get_excel_bytes_multi_sheets(sheets, file_name):
    """
    Giữ một file gần nhất mỗi báo cáo nhiều sheet trong phiên.
    sheets: list[(sheet_name, df)]. Format float '0.00' mọi sheet.
    """
    cache = st.session_state.setdefault('_excel_download_cache', {})
    key = file_name
    cached = cache.get(key)
    if cached is not None:
        cached_sheets = cached.get('sheets', [])
        if len(cached_sheets) == len(sheets) and all(
            name == c_name and df.equals(c_df)
            for (name, df), (c_name, c_df) in zip(sheets, cached_sheets)
        ):
            return cached['data']

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        for sheet_name, df_sheet in sheets:
            df_sheet.to_excel(writer, sheet_name=sheet_name, index=False)
            ws = writer.sheets[sheet_name]
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, float):
                        cell.number_format = '0.00'
    data = buffer.getvalue()
    cache[key] = {
        'sheets': [(name, df.copy()) for name, df in sheets],
        'data': data,
    }
    return data


def render_excel_download_multi_sheets(sheets, file_name, button_label):
    """1 nút download cho workbook nhiều sheet."""
    st.download_button(
        label=button_label,
        data=get_excel_bytes_multi_sheets(sheets, file_name),
        file_name=file_name,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


def _format_excel_cell(cell, val, fmt_type, is_bold, normal_font, bold_font, border):
    cell.font = bold_font if is_bold else normal_font
    cell.border = border
    if val is None or (isinstance(val, float) and pd.isna(val)):
        cell.value = ""
        return

    if fmt_type == "float":
        try:
            cell.value = float(val)
            cell.number_format = '0.00'
        except (ValueError, TypeError):
            cell.value = 0.0
            cell.number_format = '0.00'
    elif fmt_type == "pct":
        try:
            # Giá trị trong df là 0-100 (vd 20.0). Trong Excel format 0.00%, giá trị số là 0.20
            cell.value = float(val) / 100.0
            cell.number_format = '0.00%'
        except (ValueError, TypeError):
            cell.value = 0.0
            cell.number_format = '0.00%'
    elif fmt_type == "currency_int":
        try:
            cell.value = int(round(float(val)))
            cell.number_format = '#,##0'
        except (ValueError, TypeError):
            cell.value = 0
            cell.number_format = '#,##0'
    elif fmt_type == "currency_float":
        try:
            cell.value = float(val)
            cell.number_format = '#,##0.00'
        except (ValueError, TypeError):
            cell.value = 0.0
            cell.number_format = '#,##0.00'
    else:
        cell.value = str(val)


def build_report_5_excel_bytes(sheets, file_name):
    """
    Tạo file Excel chuyên biệt cho Báo cáo 5 với 2 hàng header:
    - Hàng 1: Nhóm cột (merge các cột con tương ứng; cột đơn merge hàng 1-2).
    - Hàng 2: Tên cột con.
    - Cố định 2 hàng header và 2 cột đầu (freeze_panes = 'C3').
    - Bật AutoFilter.
    - Dòng TỔNG CỘNG in đậm.
    """
    # Mỗi item có thể là (tên_sheet, dataframe) hoặc (tên_sheet, dataframe, include_costs).
    sheet_items = [
        (item[0], item[1], item[2] if len(item) >= 3 else True)
        for item in sheets
    ]
    cache = st.session_state.setdefault('_excel_download_cache', {})
    key = file_name
    cached = cache.get(key)
    if cached is not None:
        cached_sheets = cached.get('sheets', [])
        if len(cached_sheets) == len(sheet_items) and all(len(item) == 3 for item in cached_sheets) and all(
            name == c_name and include_costs == c_include_costs and df.equals(c_df)
            for (name, df, include_costs), (c_name, c_df, c_include_costs) in zip(sheet_items, cached_sheets)
        ):
            return cached['data']

    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from report_5_schema import get_report_5_column_specs, FIELD_TIME
    from report_calculations import aggregate_report_5_rows

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    header_fill = PatternFill(start_color="E8F0FE", end_color="E8F0FE", fill_type="solid")
    header_group_fill = PatternFill(start_color="D2E3FC", end_color="D2E3FC", fill_type="solid")
    bold_font = Font(name="Calibri", size=11, bold=True)
    normal_font = Font(name="Calibri", size=11)
    thin_border = Border(
        left=Side(style='thin', color='D3D3D3'),
        right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'),
        bottom=Side(style='thin', color='D3D3D3')
    )

    for sheet_name, df_table, include_costs in sheet_items:
        specs = get_report_5_column_specs(include_costs=include_costs)
        flat_cols = []
        for grp_name, children in specs:
            for child_name, field_key, fmt_type, width in children:
                flat_cols.append((child_name, field_key, fmt_type, width, grp_name))
        ws = wb.create_sheet(title=sheet_name)
        ws.views.sheetView[0].showGridLines = True

        col_idx = 1
        for grp_name, children in specs:
            start_col = col_idx
            num_children = len(children)
            end_col = start_col + num_children - 1

            if grp_name is None:
                for child_name, _, _, _ in children:
                    c1 = ws.cell(row=1, column=col_idx, value=child_name)
                    c1.font = bold_font
                    c1.fill = header_fill
                    c1.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    c1.border = thin_border

                    c2 = ws.cell(row=2, column=col_idx, value="")
                    c2.border = thin_border
                    ws.merge_cells(start_row=1, start_column=col_idx, end_row=2, end_column=col_idx)
                    col_idx += 1
            else:
                for c_in_grp in range(start_col, end_col + 1):
                    cell = ws.cell(row=1, column=c_in_grp)
                    cell.border = thin_border
                    cell.fill = header_group_fill
                c_grp = ws.cell(row=1, column=start_col, value=grp_name)
                c_grp.font = bold_font
                c_grp.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)

                for child_name, _, _, _ in children:
                    c_child = ws.cell(row=2, column=col_idx, value=child_name)
                    c_child.font = bold_font
                    c_child.fill = header_fill
                    c_child.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    c_child.border = thin_border
                    col_idx += 1

        current_row = 3
        if df_table is not None and not df_table.empty:
            for _, row_data in df_table.iterrows():
                for c_idx, (_, field_key, fmt_type, _, _) in enumerate(flat_cols, start=1):
                    cell = ws.cell(row=current_row, column=c_idx)
                    val = row_data.get(field_key)
                    _format_excel_cell(cell, val, fmt_type, is_bold=False, normal_font=normal_font, bold_font=bold_font, border=thin_border)
                current_row += 1

            time_val = df_table[FIELD_TIME].iloc[0] if len(df_table) > 0 else ""
            tot_row = aggregate_report_5_rows(df_table, time_val, "TỔNG CỘNG")
            for c_idx, (_, field_key, fmt_type, _, _) in enumerate(flat_cols, start=1):
                cell = ws.cell(row=current_row, column=c_idx)
                val = tot_row.get(field_key)
                _format_excel_cell(cell, val, fmt_type, is_bold=True, normal_font=normal_font, bold_font=bold_font, border=thin_border)
            current_row += 1

        for c_idx, (_, _, _, width, _) in enumerate(flat_cols, start=1):
            col_letter = get_column_letter(c_idx)
            ws.column_dimensions[col_letter].width = max(12, int(width / 7.5))

        ws.freeze_panes = "C3"
        total_cols = len(flat_cols)
        data_end_row = max(current_row - 1, 2)
        ws.auto_filter.ref = f"A2:{get_column_letter(total_cols)}{data_end_row}"

    buffer = io.BytesIO()
    wb.save(buffer)
    data = buffer.getvalue()
    cache[key] = {
        'sheets': [(name, df.copy(), include_costs) for name, df, include_costs in sheet_items],
        'data': data,
    }
    return data


def render_report_5_excel_download(sheets, file_name, button_label):
    """1 nút download riêng cho Báo cáo 5 với định dạng 2 hàng header và number formatting."""
    st.download_button(
        label=button_label,
        data=build_report_5_excel_bytes(sheets, file_name),
        file_name=file_name,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key="btn_download_excel_report_5",
    )


def build_report_6_excel_bytes(sheets, file_name):
    """Tạo workbook BC6 gồm nhiều sheet, mỗi sheet có header TC hai tầng."""
    cache = st.session_state.setdefault('_excel_download_cache', {})
    cached = cache.get(file_name)
    if cached is not None:
        cached_sheets = cached.get('sheets', [])
        if len(cached_sheets) == len(sheets) and all(
            name == cached_name and df.equals(cached_df)
            for (name, df), (cached_name, cached_df) in zip(sheets, cached_sheets)
        ):
            return cached['data']

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from report_6_schema import get_report_6_column_specs, FIELD_TIME
    from report_calculations import aggregate_report_6_rows

    specs = get_report_6_column_specs()
    flat_columns = [
        (child_name, field, fmt_type, width)
        for _, children in specs for child_name, field, fmt_type, width in children
    ]
    header_fill = PatternFill(start_color='E8F0FE', end_color='E8F0FE', fill_type='solid')
    group_fill = PatternFill(start_color='D2E3FC', end_color='D2E3FC', fill_type='solid')
    bold_font = Font(name='Calibri', size=11, bold=True)
    normal_font = Font(name='Calibri', size=11)
    border = Border(
        left=Side(style='thin', color='D3D3D3'), right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'), bottom=Side(style='thin', color='D3D3D3'),
    )
    wb = Workbook()
    wb.remove(wb.active)

    for sheet_name, df_table in sheets:
        ws = wb.create_sheet(title=sheet_name)
        col_index = 1
        for group_name, children in specs:
            start = col_index
            end = start + len(children) - 1
            if group_name is None:
                for child_name, _, _, _ in children:
                    cell = ws.cell(1, col_index, child_name)
                    cell.font, cell.fill, cell.border = bold_font, header_fill, border
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                    ws.cell(2, col_index).border = border
                    ws.merge_cells(start_row=1, start_column=col_index, end_row=2, end_column=col_index)
                    col_index += 1
            else:
                for current in range(start, end + 1):
                    ws.cell(1, current).fill = group_fill
                    ws.cell(1, current).border = border
                group_cell = ws.cell(1, start, group_name)
                group_cell.font = bold_font
                group_cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
                for child_name, _, _, _ in children:
                    cell = ws.cell(2, col_index, child_name)
                    cell.font, cell.fill, cell.border = bold_font, header_fill, border
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                    col_index += 1

        output = df_table.copy()
        if not output.empty:
            time_val = output[FIELD_TIME].iloc[0] if FIELD_TIME in output else ''
            output = pd.concat(
                [output, pd.DataFrame([aggregate_report_6_rows(output, time_val)])],
                ignore_index=True,
            )
        for row_index, (_, row_data) in enumerate(output.iterrows(), start=3):
            is_total = str(row_data.get('Tỉnh/Thành phố', '')).startswith('📊 TỔNG')
            for column_index, (_, field, fmt_type, _) in enumerate(flat_columns, start=1):
                _format_excel_cell(
                    ws.cell(row_index, column_index), row_data.get(field), fmt_type,
                    is_total, normal_font, bold_font, border,
                )
        for column_index, (_, _, _, width) in enumerate(flat_columns, start=1):
            ws.column_dimensions[get_column_letter(column_index)].width = max(12, int(width / 7.5))
        ws.freeze_panes = 'C3'
        ws.auto_filter.ref = f"A2:{get_column_letter(len(flat_columns))}{max(2, len(output) + 2)}"

    buffer = io.BytesIO()
    wb.save(buffer)
    data = buffer.getvalue()
    cache[file_name] = {'sheets': [(name, df.copy()) for name, df in sheets], 'data': data}
    return data


def render_report_6_excel_download(sheets, file_name, button_label):
    st.download_button(
        label=button_label,
        data=build_report_6_excel_bytes(sheets, file_name),
        file_name=file_name,
        mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        key='btn_download_excel_report_6',
    )




def assign_dot_manual_to_first_row(df, dot_manual_df, columns, type_map=None):
    """
    Phân bổ giá trị nhập tay theo đợt vào dòng đầu tiên của mỗi đợt trong DataFrame chính.
    Loại bỏ trùng lặp logic giữa render_report_1 và render_report_2.

    Args:
        df: DataFrame chính (sẽ bị thay đổi inplace).
        dot_manual_df: DataFrame manual inputs theo đợt (có cột ĐỢT HỌC THỬ).
        columns: list tên cột cần gán.
        type_map: dict {col_name: callable} để ép kiểu. Mặc định: int.
                  Ví dụ: {'Data trùng bình quân ...': float}

    Returns:
        df: DataFrame đã được gán giá trị.
    """
    if type_map is None:
        type_map = {}
    default_values = {}
    for col in columns:
        cast_fn = type_map.get(col, int)
        # Xác định giá trị mặc định dựa trên kiểu
        default_values[col] = 0.0 if cast_fn is float else 0
        df[col] = default_values[col]

    for _, row in dot_manual_df.iterrows():
        dot_mask = df[DOT_COLUMN] == row[DOT_COLUMN]
        if dot_mask.any():
            first_idx = df[dot_mask].index[0]
            for col in columns:
                cast_fn = type_map.get(col, int)
                df.at[first_idx, col] = cast_fn(row[col])

    return df
